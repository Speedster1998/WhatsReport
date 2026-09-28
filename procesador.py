import re
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border as ExcelBorder, Side as ExcelSide

def procesar_reporte(ruta_txt: str, ruta_salida_excel: str):
    msg_pattern = re.compile(
        r'\[(\d{1,2}/\d{1,2}/\d{2,4}),\s*(\d{1,2}:\d{2}:\d{2}\s*[APMapm\.]+)\]\s*([^:]+):\s*(.*?)(?=\n\[\d{1,2}/\d{1,2}/\d{2,4}|\Z)', 
        re.DOTALL
    )

    datos_extraidos = []

    with open(ruta_txt, 'r', encoding='utf-8') as file:
        texto_completo = file.read()

    for match in msg_pattern.finditer(texto_completo):
        fecha_raw, hora_raw, remitente, contenido = match.groups()
        
        if "MONITOR" in contenido.upper() and ("CONEXI" in contenido.upper() or "CÁMARAS" in contenido.upper()):
            partes_fecha = fecha_raw.split('/')
            try:
                mes_num = int(partes_fecha[0])
                dia_num = int(partes_fecha[1])
                anio_num = int(partes_fecha[2])
                if anio_num < 100:
                    anio_num += 2000
                fecha_obj = datetime(anio_num, mes_num, dia_num).date()
            except Exception:
                fecha_obj = fecha_raw
            
            hora_obj = None
            for fmt in ("%I:%M:%S %p", "%I:%M:%S %p.", "%H:%M:%S"):
                try:
                    hora_limpia = hora_raw.replace('.', '').strip()
                    hora_obj = datetime.strptime(hora_limpia, fmt)
                    break
                except ValueError:
                    pass
            
            hora_formateada = hora_obj.strftime('%H:%M:%S') if hora_obj else hora_raw
            
            resp_match = re.search(r'Responsable\s*:\s*(.+)', contenido, re.IGNORECASE)
            responsable = resp_match.group(1).split('\n')[0].strip() if resp_match else remitente.strip()
            
            bloque_conexion, bloque_obs = "", ""
            if re.search(r'Conexi[oó]n de c[aá]maras', contenido, re.IGNORECASE):
                partes = re.split(r'Conexi[oó]n de c[aá]maras', contenido, flags=re.IGNORECASE)
                resto = partes[1]
                if re.search(r'Observaciones[/\s]*Novedades', resto, re.IGNORECASE):
                    partes_obs = re.split(r'Observaciones[/\s]*Novedades', resto, flags=re.IGNORECASE)
                    bloque_conexion = partes_obs[0]
                    bloque_obs = partes_obs[1] if len(partes_obs) > 1 else ""
                else:
                    bloque_conexion = resto

            obs_dict = {}
            for linea_obs in bloque_obs.split('\n'):
                linea_obs = linea_obs.strip()
                match_o = re.search(r'^[*\s\.]*([A-Za-z0-9]+)\s*:\s*(.*)', linea_obs)
                if match_o:
                    b_code = match_o.group(1).strip().upper()
                    o_texto = match_o.group(2).strip()
                    if o_texto and o_texto != "-":
                        obs_dict[b_code] = o_texto

            for linea_cam in bloque_conexion.split('\n'):
                linea_cam = linea_cam.strip()
                match_c = re.search(r'^[*\s]*([A-Za-z0-9]+)\s*[:\s]\s*(.+)', linea_cam)
                if match_c:
                    base = match_c.group(1).strip().upper()
                    estado = match_c.group(2).strip()

                    if base in ["TODO", "SIN", "CONEXIÓN", "CONEXION", "RESPONSABLE"]:
                        continue

                    if "TODO OK" in estado.upper():
                        conexion = "Todo ok"
                        observacion = obs_dict.get(base, "")
                    else:
                        conexion = "Sin visión" if "SIN VISI" in estado.upper() else estado
                        observacion = estado if "Sin visión de" in estado else obs_dict.get(base, estado)

                    datos_extraidos.append({
                        "Fecha": fecha_obj,
                        "Hora": hora_formateada,
                        "Base": base,
                        "Conexión cámaras": conexion,
                        "Responsable": responsable,
                        "Observaciones": observacion,
                        "Acciones tomadas": ""
                    })

    if not datos_extraidos:
        return 0, None

    df = pd.DataFrame(datos_extraidos)
    with pd.ExcelWriter(ruta_salida_excel, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Reporte Diario")
        
    wb = load_workbook(ruta_salida_excel)
    ws = wb["Reporte Diario"]
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    red_font = Font(color="FF0000")
    center_align = Alignment(horizontal="center", vertical="center")
    
    linea = ExcelSide(style='thin', color='8EA9DB')
    borde_medio = ExcelBorder(top=linea, bottom=linea)
    borde_izquierdo = ExcelBorder(left=linea, top=linea, bottom=linea)
    borde_derecho = ExcelBorder(right=linea, top=linea, bottom=linea)
    
    anchos = {'A': 12, 'B': 13, 'C': 10, 'D': 18, 'E': 22, 'F': 45, 'G': 20}
    for col, width in anchos.items():
        ws.column_dimensions[col].width = width
        
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        if cell.column == 1:
            cell.border = ExcelBorder(left=linea, top=linea, bottom=linea)
        elif cell.column == ws.max_column:
            cell.border = ExcelBorder(right=linea, top=linea, bottom=linea)
        else:
            cell.border = borde_medio
        
    ws.auto_filter.ref = ws.dimensions
    
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.column == 1:
                cell.border = borde_izquierdo
                cell.number_format = 'dd/mm/yyyy'
            elif cell.column == ws.max_column:
                cell.border = borde_derecho
            else:
                cell.border = borde_medio
            
            if cell.column == 2 and isinstance(cell.value, str):
                try:
                    t_obj = datetime.strptime(cell.value, '%H:%M:%S').time()
                    cell.value = t_obj
                    cell.number_format = 'hh:mm:ss'
                except ValueError:
                    pass

            if cell.column in (3, 5):
                cell.alignment = center_align

            if cell.column == 4 and ("Sin visión" in str(cell.value) or "Sin conexión" in str(cell.value)):
                cell.font = red_font
                
    wb.save(ruta_salida_excel)
    return len(datos_extraidos), ruta_salida_excel